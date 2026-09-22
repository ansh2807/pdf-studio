#!/usr/bin/env python3
"""
Tests for engine.py's spacing-delimited report/invoice reconstruction
(_reconstruct_report, _num_or_text, used by pdf-to-excel/pdf-to-word for
scanned-look invoices and sales summaries with no ruled table borders).

Builds REAL PDFs with precisely positioned text (not mocked word lists) and
runs them through the actual PyMuPDF word extraction + reconstruction, so
this exercises the genuine code path, not an approximation of it.

Locks in two real bugs found this way:
  1. A percentage cell ("18%") was silently converted to the bare number 18,
     losing the percent sign - actively misleading data (a rate becomes
     indistinguishable from a raw quantity/amount in the exported cell).
  2. A short, entirely realistic 3-word column header ("Product Qty Amount")
     was misclassified as a title/note line instead of a header ROW, because
     of an unrelated `> 3` word-count threshold that only happened to pass
     for the 4-word header used in earlier ad-hoc testing.

Run:  python engine/tests/test_report_reconstruction.py
"""
from __future__ import annotations
import os, sys, json, subprocess
import fitz

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
ENGINE = os.path.join(ROOT, "engine.py")
sys.path.insert(0, ROOT)
import engine as E  # noqa: E402

WORK = os.path.join(HERE, "proof")


def build_invoice(path, rows_data, cols, header=None, title_lines=None, pages=1):
    d = fitz.open()
    for pg in range(pages):
        p = d.new_page(width=595, height=400)
        y = 40
        if title_lines:
            for t in title_lines:
                p.insert_text((40, y), t, fontsize=11)
                y += 15
        y += 15
        if header:
            for i, h in enumerate(header):
                p.insert_text((cols[i], y), h, fontsize=9)
            y += 20
        for row in rows_data:
            for i, val in enumerate(row):
                p.insert_text((cols[i], y), str(val), fontsize=9)
            y += 16
    d.save(path)
    d.close()


RESULTS = []


def chk(name, cond, detail=""):
    RESULTS.append((name, bool(cond)))
    print(f"  {'PASS' if cond else 'FAIL'}  {name}  {detail}")


def main():
    os.makedirs(WORK, exist_ok=True)
    cols4 = [40, 260, 340, 420]

    # ---------- bug 1: percentage preservation ----------
    p1 = os.path.join(WORK, "report_pct.pdf")
    build_invoice(p1, [
        ("Widget A", "10", "5.00", "50.00"),
        ("Gadget Deluxe", "3", "1,250.75", "3,752.25"),
        ("Service Fee", "-", "-", "18%"),
        ("Discount Item", "2", "-100.00", "-200.00"),
    ], cols4, header=["DESCRIPTION", "QTY", "RATE", "AMOUNT"],
       title_lines=["ACME TRADERS PVT LTD", "123 Market Street"])
    doc = fitz.open(p1)
    rows, ncols, ndata = E._reconstruct_report(doc)
    doc.close()
    data_rows = [r for r in rows if r["kind"] == "row"]
    pct_row = next((r for r in data_rows if r["cells"][0] == "Service Fee"), None)
    chk("percentage cell keeps its % sign, not a bare misleading number",
        pct_row is not None and pct_row["cells"][3] == "18%",
        pct_row["cells"] if pct_row else None)
    comma_row = next((r for r in data_rows if r["cells"][0] == "Gadget Deluxe"), None)
    chk("comma-formatted currency parses as a real, usable, correctly-valued number",
        comma_row is not None and comma_row["cells"][2] == 1250.75 and comma_row["cells"][3] == 3752.25,
        comma_row["cells"] if comma_row else None)
    widget_row = next((r for r in data_rows if r["cells"][0] == "Widget A"), None)
    chk("plain currency value converts to a real number (summable in Excel)",
        widget_row is not None and widget_row["cells"][3] == 50, widget_row)
    neg_row = next((r for r in data_rows if r["cells"][0] == "Discount Item"), None)
    chk("negative amounts are preserved as real negative numbers",
        neg_row is not None and neg_row["cells"][3] == -200, neg_row)
    dash_row = pct_row
    chk("dash placeholders become blank cells (None), not the literal text '-'",
        dash_row is not None and dash_row["cells"][1] is None and dash_row["cells"][2] is None,
        dash_row["cells"] if dash_row else None)
    header_row = next((r for r in data_rows if r["cells"] == ["DESCRIPTION", "QTY", "RATE", "AMOUNT"]), None)
    chk("4-word header row is correctly identified as a header, not a title",
        header_row is not None, header_row)
    titles = [r["text"] for r in rows if r["kind"] == "title"]
    chk("title/address lines are captured as titles, not merged into data",
        "ACME TRADERS PVT LTD" in titles and "123 Market Street" in titles, titles)

    # ---------- bug 2: short (3-word) header must not be misclassified ----------
    p2 = os.path.join(WORK, "report_short_header.pdf")
    cols3 = [40, 300, 400]
    build_invoice(p2, [("Item One", "5", "100")], cols3,
                 header=["Product", "Qty", "Amount"],
                 title_lines=["ITEM WISE SALES SUMMARY"])
    doc2 = fitz.open(p2)
    rows2, _, _ = E._reconstruct_report(doc2)
    doc2.close()
    header2 = next((r for r in rows2 if r["kind"] == "row" and r["cells"] == ["Product", "Qty", "Amount"]), None)
    chk("3-word header ('Product Qty Amount') is correctly identified as a header row",
        header2 is not None, [r for r in rows2])
    title2 = next((r for r in rows2 if r["kind"] == "title" and "ITEM WISE SALES SUMMARY" in r["text"]), None)
    chk("a section title containing ONE keyword ('ITEM') stays a title, not a false-positive header",
        title2 is not None, title2)

    # ---------- multi-page: repeating company-name title dedups; data doesn't ----------
    p3 = os.path.join(WORK, "report_multipage.pdf")
    d = fitz.open()
    for pg in range(2):
        pg_obj = d.new_page(width=595, height=250)
        pg_obj.insert_text((40, 40), "ACME TRADERS PVT LTD", fontsize=11)
        for i, h in enumerate(["Product", "Qty", "Amount"]):
            pg_obj.insert_text((cols3[i], 70), h, fontsize=9)
        for i, v in enumerate([f"Item Page{pg}", "1", "10"]):
            pg_obj.insert_text((cols3[i], 90), v, fontsize=9)
    d.save(p3); d.close()
    doc3 = fitz.open(p3)
    rows3, _, _ = E._reconstruct_report(doc3)
    doc3.close()
    title_count = sum(1 for r in rows3 if r["kind"] == "title" and r["text"] == "ACME TRADERS PVT LTD")
    chk("a repeating title/company-name line dedups across pages (appears once)",
        title_count == 1, title_count)
    data_items = [r["cells"][0] for r in rows3 if r["kind"] == "row" and str(r["cells"][0]).startswith("Item Page")]
    chk("actual per-page DATA rows are NOT deduplicated (both pages' rows survive)",
        data_items == ["Item Page0", "Item Page1"], data_items)

    # ---------- end-to-end: real `pdf-to-word` CLI command, not just the ----------
    # ---------- internal function. cmd_pdf_to_word only takes the report   ----------
    # ---------- path when report_ndata >= 5, a threshold the 4-row fixture ----------
    # ---------- above never exercises - so this is the only test that      ----------
    # ---------- proves the percentage/header fixes are reachable in real   ----------
    # ---------- usage of the command the server actually calls.            ----------
    p4 = os.path.join(WORK, "full_invoice_5row.pdf")
    build_invoice(p4, [
        ("Widget A", "10", "5.00", "50.00"),
        ("Gadget B Deluxe", "3", "1,250.75", "3,752.25"),
        ("Service Fee", "-", "-", "18%"),
        ("Discount Item", "2", "-100.00", "-200.00"),
        ("Consulting Hours", "8", "75.00", "600.00"),
    ], cols4, header=["DESCRIPTION", "QTY", "RATE", "AMOUNT"],
       title_lines=["ACME TRADERS PVT LTD"])
    out4 = os.path.join(WORK, "full_invoice_5row.docx")
    proc = subprocess.run(
        [sys.executable, ENGINE, "pdf-to-word", "--input", p4, "--output", out4],
        capture_output=True, text=True)
    try:
        result = json.loads(proc.stdout.strip().splitlines()[-1])
    except Exception:
        result = {}
    chk("CLI command succeeds and takes the report-reconstruction path (not generic pdf2docx)",
        result.get("ok") and result.get("fidelity") == "report" and result.get("dataRows") == 5,
        (proc.stdout, proc.stderr))

    from docx import Document
    doc4 = Document(out4)
    table = doc4.tables[0]
    cells = [[c.text for c in row.cells] for row in table.rows]
    chk("real docx: header splits into 4 distinct cells, not merged into the title",
        cells[0] == ["DESCRIPTION", "QTY", "RATE", "AMOUNT"], cells[0])
    chk("real docx: percentage cell keeps its % sign through the actual CLI command",
        cells[3] == ["Service Fee", "", "", "18%"], cells[3])
    chk("real docx: comma-formatted and negative currency cells are correct",
        cells[2][2] == "1250.75" and cells[4][3] == "-200", (cells[2], cells[4]))
    title_para = next((p for p in doc4.paragraphs if p.text.strip()), None)
    chk("real docx: company title is a bolded paragraph, not folded into the table",
        title_para is not None and title_para.text == "ACME TRADERS PVT LTD"
        and title_para.runs and title_para.runs[0].font.bold,
        title_para.text if title_para else None)

    # ---------- same 5-row fixture through the real `pdf-to-excel` CLI ----------
    # ---------- command - a separate code path (_report_to_sheet) that   ----------
    # ---------- also consumes _reconstruct_report's output directly.     ----------
    out5 = os.path.join(WORK, "full_invoice_5row.xlsx")
    proc2 = subprocess.run(
        [sys.executable, ENGINE, "pdf-to-excel", "--input", p4, "--output", out5],
        capture_output=True, text=True)
    try:
        result2 = json.loads(proc2.stdout.strip().splitlines()[-1])
    except Exception:
        result2 = {}
    chk("pdf-to-excel CLI succeeds and takes the report-reconstruction path",
        result2.get("ok") and result2.get("mode") == "report" and result2.get("rows") == 7,
        (proc2.stdout, proc2.stderr))

    import openpyxl
    wb = openpyxl.load_workbook(out5)
    ws = wb["Report"]
    xrows = list(ws.iter_rows(values_only=True))
    chk("real xlsx: header splits into 4 distinct cells",
        xrows[1] == ("DESCRIPTION", "QTY", "RATE", "AMOUNT"), xrows[1])
    chk("real xlsx: percentage cell keeps its % sign through the actual CLI command",
        xrows[4] == ("Service Fee", None, None, "18%"), xrows[4])
    chk("real xlsx: negative/comma amounts are real numbers, summable in Excel",
        xrows[3][2] == 1250.75 and xrows[5][3] == -200, (xrows[3], xrows[5]))
    chk("real xlsx: company title occupies its own row and is bolded",
        xrows[0] == ("ACME TRADERS PVT LTD", None, None, None) and ws["A1"].font.bold,
        xrows[0])

    passed = sum(1 for _, ok in RESULTS if ok)
    total = len(RESULTS)
    print(f"\nREPORT-RECONSTRUCTION: {passed}/{total} passed")
    print("RESULT:", "PASS" if passed == total else "FAIL")
    return 0 if passed == total else 1


if __name__ == "__main__":
    sys.exit(main())
